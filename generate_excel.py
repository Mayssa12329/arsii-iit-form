import json
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load data
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
json_path = os.path.join(BASE_DIR, "data", "adhesions.json")
with open(json_path, "r", encoding="utf-8") as f:
    adhesions = json.load(f)

# Sort by ID ascending (or descending, let's keep ID 1 to 70 for clean chronological order or preserve user's list)
# Usually chronological ID 1 -> 70 is best for an official registry, or descending if most recent first.
# Let's sort by ID ascending for registry clarity.
adhesions_sorted = sorted(adhesions, key=lambda x: x.get("id", 0))

wb = openpyxl.Workbook()

# Sheet 1: Liste des Adhésions
ws1 = wb.active
ws1.title = "Adhésions"
ws1.views.sheetView[0].showGridLines = True

# Palettes
NAVY_HEADER = "1B365D"      # ARSII Deep Navy Blue
ACCENT_BLUE = "0284C7"      # Vibrant Blue
WHITE = "FFFFFF"
ZEBRA_FILL = "F8FAFC"       # Very subtle light slate
BORDER_COLOR = "E2E8F0"
GREEN_FILL = "DCFCE7"       # Soft green for email sent
GREEN_TEXT = "166534"
GRAY_TEXT = "475569"

font_title = Font(name="Segoe UI", size=15, bold=True, color="1B365D")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
font_header = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
font_data = Font(name="Segoe UI", size=10, color="1E293B")
font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
font_badge = Font(name="Segoe UI", size=9, bold=True, color=GREEN_TEXT)

header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
zebra_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
badge_fill = PatternFill(start_color=GREEN_FILL, end_color=GREEN_FILL, fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_COLOR)
data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
header_border = Border(
    left=Side(border_style="thin", color="334E68"),
    right=Side(border_style="thin", color="334E68"),
    top=Side(border_style="medium", color="0F2C59"),
    bottom=Side(border_style="medium", color="0F2C59")
)

# Title Block
ws1.merge_cells("A1:M1")
ws1["A1"] = "ARSII IIT — Registre Officiel des Adhésions"
ws1["A1"].font = font_title
ws1["A1"].alignment = Alignment(vertical="center", indent=1)
ws1.row_dimensions[1].height = 36

ws1.merge_cells("A2:M2")
total_count = len(adhesions_sorted)
ws1["A2"] = f"Total des candidatures : {total_count} inscriptions | Année universitaire 2026/2027 | Club ARSII IIT"
ws1["A2"].font = font_subtitle
ws1["A2"].alignment = Alignment(vertical="center", indent=1)
ws1.row_dimensions[2].height = 20

# Column Headers
headers = [
    ("ID", 8, "center"),
    ("Date & Heure", 18, "center"),
    ("Nom", 16, "left"),
    ("Prénom", 16, "left"),
    ("Email", 28, "left"),
    ("Téléphone", 16, "center"),
    ("Niveau", 18, "left"),
    ("Filière", 22, "left"),
    ("Année", 14, "center"),
    ("Centres d'intérêt", 38, "left"),
    ("Motivation", 42, "left"),
    ("Source", 20, "left"),
    ("Confirmation", 15, "center")
]

start_row = 4
ws1.row_dimensions[start_row].height = 28

for col_idx, (header_text, _, align) in enumerate(headers, 1):
    cell = ws1.cell(row=start_row, column=col_idx, value=header_text)
    cell.font = font_header
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = header_border

# Data Rows
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_wrap_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

for i, row in enumerate(adhesions_sorted):
    current_row = start_row + 1 + i
    ws1.row_dimensions[current_row].height = 26
    is_even = (i % 2 == 1)
    current_fill = zebra_fill if is_even else white_fill
    
    # Format date
    raw_date = row.get("created_at", "")
    formatted_date = raw_date
    if raw_date:
        try:
            dt = datetime.fromisoformat(raw_date.replace("Z", "+00:00"))
            formatted_date = dt.strftime("%Y-%m-%d %H:%M")
        except Exception:
            formatted_date = raw_date

    email_status = "Envoyé" if row.get("email_sent") == 1 else "Non envoyé"

    values = [
        (row.get("id"), align_center, font_data_bold),
        (formatted_date, align_center, font_data),
        (row.get("nom", "").strip(), align_left, font_data_bold),
        (row.get("prenom", "").strip(), align_left, font_data),
        (row.get("email", "").strip(), align_left, font_data),
        (str(row.get("telephone", "")).strip(), align_center, font_data),
        (row.get("niveau", "").strip(), align_left, font_data),
        (row.get("filiere", "").strip(), align_left, font_data),
        (row.get("annee", "").strip(), align_center, font_data),
        (row.get("interets", "").strip(), align_wrap_left, font_data),
        (row.get("motivation", "").strip().replace("\n", " "), align_wrap_left, font_data),
        (row.get("source", "").strip(), align_left, font_data),
        (email_status, align_center, font_badge)
    ]

    for col_idx, (val, alignment, font_style) in enumerate(values, 1):
        cell = ws1.cell(row=current_row, column=col_idx, value=val)
        cell.alignment = alignment
        cell.font = font_style
        cell.border = data_border
        
        # Badge styling for status
        if col_idx == 13:
            cell.fill = badge_fill
        else:
            cell.fill = current_fill

# Set column widths
for col_idx, (_, width, _) in enumerate(headers, 1):
    col_letter = get_column_letter(col_idx)
    ws1.column_dimensions[col_letter].width = width

# Enable autofilter on table
ws1.auto_filter.ref = f"A{start_row}:M{start_row + len(adhesions_sorted)}"
# Freeze panes below header
ws1.freeze_panes = f"A{start_row + 1}"


# ==========================================
# Sheet 2: Statistiques & Synthèse
# ==========================================
ws2 = wb.create_sheet(title="Statistiques & Synthèse")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:G1")
ws2["A1"] = "ARSII IIT — Tableau de Bord & Synthèse des Inscriptions"
ws2["A1"].font = font_title
ws2["A1"].alignment = Alignment(vertical="center", indent=1)
ws2.row_dimensions[1].height = 36

# KPI summary cards
kpis = [
    ("Total Inscrits", str(total_count), "0284C7"),
    ("Cycle Ingénieur", str(sum(1 for r in adhesions if "ingénieur" in r.get("niveau", "").lower())), "16A34A"),
    ("Licence", str(sum(1 for r in adhesions if "licence" in r.get("niveau", "").lower() or "license" in r.get("niveau", "").lower())), "D97706"),
    ("Emails Confirmés", f"{sum(1 for r in adhesions if r.get('email_sent') == 1)} / {total_count}", "7C3AED")
]

ws2.row_dimensions[3].height = 18
ws2.row_dimensions[4].height = 32

for idx, (label, val, color_code) in enumerate(kpis):
    col_start = 1 + idx * 2
    col_end = col_start + 1
    col_start_letter = get_column_letter(col_start)
    col_end_letter = get_column_letter(col_end)
    
    ws2.merge_cells(f"{col_start_letter}3:{col_end_letter}3")
    ws2.merge_cells(f"{col_start_letter}4:{col_end_letter}4")
    
    c_label = ws2[f"{col_start_letter}3"]
    c_label.value = label.upper()
    c_label.font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    c_label.alignment = Alignment(horizontal="center", vertical="center")
    c_label.fill = PatternFill(start_color="F1F5F9", fill_type="solid")
    
    c_val = ws2[f"{col_start_letter}4"]
    c_val.value = val
    c_val.font = Font(name="Segoe UI", size=18, bold=True, color=color_code)
    c_val.alignment = Alignment(horizontal="center", vertical="center")
    c_val.fill = PatternFill(start_color="F8FAFC", fill_type="solid")
    
    card_border = Border(
        left=Side(style="medium", color=color_code),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    for r in range(3, 5):
        for c in range(col_start, col_end + 1):
            ws2.cell(row=r, column=c).border = card_border

# Function to write a mini summary table
def write_summary_table(ws, title, data_dict, start_r, start_c):
    font_tbl_hdr = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    fill_tbl_hdr = PatternFill(start_color=NAVY_HEADER, fill_type="solid")
    
    ws.merge_cells(start_row=start_r, start_column=start_c, end_row=start_r, end_column=start_c+2)
    tcell = ws.cell(row=start_r, column=start_c, value=title)
    tcell.font = Font(name="Segoe UI", size=11, bold=True, color=NAVY_HEADER)
    tcell.alignment = Alignment(vertical="center")
    ws.row_dimensions[start_r].height = 24
    
    # Headers
    h_r = start_r + 1
    ws.row_dimensions[h_r].height = 22
    for ci, hname in enumerate(["Catégorie", "Nombre", "%"]):
        hc = ws.cell(row=h_r, column=start_c + ci, value=hname)
        hc.font = font_tbl_hdr
        hc.fill = fill_tbl_hdr
        hc.alignment = Alignment(horizontal="center", vertical="center")
        hc.border = header_border
        
    cur_r = h_r + 1
    total_items = sum(data_dict.values())
    for k, v in data_dict.items():
        pct = (v / total_items * 100) if total_items > 0 else 0
        c1 = ws.cell(row=cur_r, column=start_c, value=k)
        c2 = ws.cell(row=cur_r, column=start_c+1, value=v)
        c3 = ws.cell(row=cur_r, column=start_c+2, value=f"{pct:.1f}%")
        
        c1.font = font_data
        c2.font = font_data_bold
        c3.font = font_data
        
        c1.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        
        c1.border = data_border
        c2.border = data_border
        c3.border = data_border
        
        c1.fill = white_fill
        c2.fill = white_fill
        c3.fill = white_fill
        ws.row_dimensions[cur_r].height = 20
        cur_r += 1
        
    return cur_r

# Aggregations
niveaux = {}
for r in adhesions:
    niv = r.get("niveau", "Non précisé").strip().capitalize()
    niveaux[niv] = niveaux.get(niv, 0) + 1

annees = {}
for r in adhesions:
    an = r.get("annee", "Non précisé").strip()
    annees[an] = annees.get(an, 0) + 1

sources = {}
for r in adhesions:
    src = r.get("source", "Non précisé").strip()
    sources[src] = sources.get(src, 0) + 1

interets_count = {}
for r in adhesions:
    raw_int = r.get("interets", "")
    for item in raw_int.split(","):
        cleaned = item.strip()
        if cleaned:
            interets_count[cleaned] = interets_count.get(cleaned, 0) + 1

# Sort interests descending
interets_count = dict(sorted(interets_count.items(), key=lambda x: x[1], reverse=True))

r1 = write_summary_table(ws2, "Répartition par Niveau", niveaux, 6, 1)
r2 = write_summary_table(ws2, "Répartition par Année", annees, r1 + 1, 1)

r3 = write_summary_table(ws2, "Source de Découverte", sources, 6, 5)
r4 = write_summary_table(ws2, "Domaines d'Intérêt", interets_count, r3 + 1, 5)

ws2.column_dimensions["A"].width = 24
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 12
ws2.column_dimensions["D"].width = 4
ws2.column_dimensions["E"].width = 30
ws2.column_dimensions["F"].width = 12
ws2.column_dimensions["G"].width = 12
ws2.column_dimensions["H"].width = 4

# Save workbook
output_path1 = os.path.join(BASE_DIR, "adhesions_arsii.xlsx")
output_path2 = os.path.join(BASE_DIR, "data", "adhesions_arsii.xlsx")

saved_paths = []
try:
    wb.save(output_path1)
    saved_paths.append(output_path1)
except PermissionError:
    print(f"WARNING: '{output_path1}' est ouvert dans Excel. Fermez-le pour le mettre a jour.")
except Exception as e:
    print(f"ERROR saving {output_path1}: {e}")

try:
    wb.save(output_path2)
    saved_paths.append(output_path2)
except Exception as e:
    print(f"ERROR saving {output_path2}: {e}")

if saved_paths:
    print("SUCCESS: Excel generated at", " and ".join(saved_paths))
